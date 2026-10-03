import json
import logging
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from config import settings
from ai.schemas import LectureAnalysisResult, ActionItemsDigest, IllustratedLongread, ActionItem, LectureSection
from ai.prompts import SYSTEM_INSTRUCTION, create_lecture_analysis_prompt

logger = logging.getLogger(__name__)


def extract_slide_timecodes(slide_paths: List[Union[str, Path]]) -> List[Dict[str, Any]]:
    """
    Extracts relative chronological timecodes (MM:SS) for slide images.
    Supports timestamps embedded in filename:
    - 'slide_001_2026-09-15_10-09-30.png' -> calculates offset from the first slide/start of capture
    - 'slide_001_14m20s.png' -> '14:20'
    - 'slide_001_120s.png' -> '02:00'
    Fallback -> '00:00'
    """
    results: List[Dict[str, Any]] = []
    if not slide_paths:
        return results

    valid_paths = [Path(p) for p in slide_paths if Path(p).exists()]
    if not valid_paths:
        return results

    ts_pattern = re.compile(r"(\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2})")
    parsed_dts = []
    for p in valid_paths:
        match = ts_pattern.search(p.stem)
        if match:
            try:
                dt = datetime.strptime(match.group(1), "%Y-%m-%d_%H-%M-%S")
                parsed_dts.append((p, dt))
            except ValueError:
                parsed_dts.append((p, None))
        else:
            parsed_dts.append((p, None))

    valid_dts = [dt for _, dt in parsed_dts if dt is not None]
    if valid_dts:
        base_dt = min(valid_dts)
        for p, dt in parsed_dts:
            if dt is not None:
                offset_sec = max(0.0, (dt - base_dt).total_seconds())
                m = int(offset_sec // 60)
                s = int(offset_sec % 60)
                tc = f"{m:02d}:{s:02d}"
            else:
                tc = "00:00"
            results.append({"path": p, "filename": p.name, "timecode": tc})
        return results

    rel_pattern_ms = re.compile(r"(\d+)m(\d+)s")
    rel_pattern_s = re.compile(r"_(\d+)s(?:$|\b)")

    for p in valid_paths:
        stem = p.stem
        ms_match = rel_pattern_ms.search(stem)
        s_match = rel_pattern_s.search(stem)
        if ms_match:
            tc = f"{int(ms_match.group(1)):02d}:{int(ms_match.group(2)):02d}"
        elif s_match:
            sec = int(s_match.group(1))
            tc = f"{sec // 60:02d}:{sec % 60:02d}"
        else:
            tc = "00:00"
        results.append({"path": p, "filename": p.name, "timecode": tc})

    return results


def _upload_file_with_retry(client, file_path: Path, max_retries: int = 3):
    """Uploads a file to Google Files API with exponential backoff for rate limits/transient errors."""
    last_err = None
    for attempt in range(1, max_retries + 1):
        try:
            return client.files.upload(file=str(file_path))
        except Exception as e:
            last_err = e
            err_msg = str(e).lower()
            is_rate_limit = any(
                term in err_msg
                for term in ["429", "503", "resource_exhausted", "quota", "rate limit", "temporarily unavailable"]
            )
            if is_rate_limit and attempt < max_retries:
                sleep_sec = 2 * (2 ** (attempt - 1))
                logger.warning(
                    f"Rate limit/transient error uploading {file_path.name} (attempt {attempt}/{max_retries}): {e}. "
                    f"Backing off for {sleep_sec}s..."
                )
                time.sleep(sleep_sec)
            else:
                logger.error(f"Failed uploading {file_path.name} on attempt {attempt}/{max_retries}: {e}")
                raise e
    raise last_err


class LectureAnalyzer:
    """
    Multimodal lecture analyzer leveraging Google Gemini Flash models.
    Supports direct upload of audio (MP3/WAV) and keyframe screenshots via Google GenAI Files API.
    """

    def __init__(self, api_key: Optional[str] = None, model: Optional[str] = None):
        self.api_key = api_key or settings.gemini_api_key
        self.model_name = model or settings.gemini_model
        self.simulation_mode = settings.simulation_mode or not bool(self.api_key)

        if not self.simulation_mode:
            try:
                from google import genai
                self.client = genai.Client(api_key=self.api_key)
            except ImportError:
                logger.warning("google-genai package not installed. Falling back to simulation mode.")
                self.simulation_mode = True
        else:
            self.client = None

    def analyze(
        self,
        audio_path: Optional[Union[str, Path]] = None,
        slide_paths: Optional[List[Union[str, Path]]] = None,
        attachment_paths: Optional[List[Union[str, Path]]] = None,
        subject_hint: str = "",
        is_audio_silent: bool = False,
    ) -> LectureAnalysisResult:
        """
        Analyzes audio recording, slide images, and supplementary instructor files to produce Action Items and Longread.
        """
        if self.simulation_mode:
            logger.info("Running LectureAnalyzer in SIMULATION mode.")
            return self._generate_simulated_result(subject_hint, slide_paths, attachment_paths)

        return self._run_gemini_analysis(
            audio_path,
            slide_paths,
            attachment_paths=attachment_paths,
            subject_hint=subject_hint,
            is_audio_silent=is_audio_silent,
        )

    def _run_gemini_analysis(
        self,
        audio_path: Optional[Union[str, Path]],
        slide_paths: Optional[List[Union[str, Path]]],
        attachment_paths: Optional[List[Union[str, Path]]] = None,
        subject_hint: str = "",
        is_audio_silent: bool = False,
    ) -> LectureAnalysisResult:
        """Executes actual multimodal API call to Gemini."""
        from google.genai import types

        uploaded_files = []
        try:
            # 1. Check audio volume level and upload audio file if present
            if audio_path and Path(audio_path).exists():
                from recorder.audio import AudioRecorder
                level = AudioRecorder.check_audio_level(Path(audio_path))
                if level.get("is_silent"):
                    is_audio_silent = True
                    logger.warning(
                        f"⚠️ Audio file {audio_path} has no speech (silent max_volume={level['max_volume']} dB)."
                    )

                logger.info(f"Uploading audio file: {audio_path}")
                audio_file = self.client.files.upload(file=str(audio_path))
                uploaded_files.append(audio_file)

            # 2. Extract illustrations, tables and text from student attachments (PPTX, PDF, PNG/JPG)
            dest_slides_dir = None
            if slide_paths and len(slide_paths) > 0:
                dest_slides_dir = Path(slide_paths[0]).parent
            elif audio_path and Path(audio_path).parent.exists():
                dest_slides_dir = Path(audio_path).parent / "slides"
                dest_slides_dir.mkdir(parents=True, exist_ok=True)

            effective_slide_paths = [Path(p) for p in (slide_paths or [])]
            extracted_tables = []
            attachment_illustrations = []
            extracted_texts = []

            if attachment_paths:
                try:
                    from ai.attachment_extractor import AttachmentExtractor
                    extractor = AttachmentExtractor(output_dir=dest_slides_dir)
                    att_result = extractor.process_attachments(
                        attachment_paths=attachment_paths,
                        existing_slide_paths=effective_slide_paths,
                        dest_slides_dir=dest_slides_dir,
                    )
                    # Register new standalone illustrations
                    for ext_img in att_result.get("extracted_images", []):
                        if ext_img not in effective_slide_paths:
                            effective_slide_paths.append(ext_img)

                    extracted_tables = att_result.get("extracted_tables", [])
                    attachment_illustrations = att_result.get("illustrations_info", [])
                    extracted_texts = att_result.get("extracted_texts", [])
                    logger.info(
                        f"AttachmentExtractor finished: "
                        f"{att_result.get('upgraded_slides_count', 0)} slides upgraded, "
                        f"{len(att_result.get('extracted_images', []))} illustrations extracted, "
                        f"{len(extracted_tables)} tables extracted."
                    )
                except Exception as e:
                    logger.warning(f"Error during attachment extraction: {e}")

            # 3. Upload slide screenshots and extracted illustrations (in parallel with rate-limit retry)
            slides_info = []
            if effective_slide_paths:
                parsed_slides = extract_slide_timecodes(effective_slide_paths)
                valid_count = len(parsed_slides)
                if valid_count > 0:
                    max_workers = min(4, valid_count)
                    logger.info(
                        f"Uploading {valid_count} slide images to Gemini Files API (concurrency: {max_workers})..."
                    )
                    uploaded_slide_results = [None] * valid_count

                    with ThreadPoolExecutor(max_workers=max_workers) as executor:
                        future_to_idx = {
                            executor.submit(_upload_file_with_retry, self.client, s["path"]): idx
                            for idx, s in enumerate(parsed_slides)
                        }
                        for future in as_completed(future_to_idx):
                            idx = future_to_idx[future]
                            s_data = parsed_slides[idx]
                            try:
                                img_file = future.result()
                                uploaded_slide_results[idx] = (img_file, s_data)
                            except Exception as e:
                                logger.error(
                                    f"Failed to upload slide {s_data['filename']} after retries: {e}"
                                )

                    for item in uploaded_slide_results:
                        if item is not None:
                            img_file, s_data = item
                            uploaded_files.append(img_file)
                            slides_info.append({
                                "filename": s_data["filename"],
                                "timecode": s_data["timecode"],
                            })

            # 4. Process supplementary text attachments (code, docx, pdf, txt)
            attachment_texts = list(extracted_texts)
            if attachment_paths:
                for att in attachment_paths:
                    att_path = Path(att)
                    if not att_path.exists():
                        continue
                    ext = att_path.suffix.lower()
                    if ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".pptx", ".ppt"]:
                        # Already processed via AttachmentExtractor
                        continue
                    elif ext == ".docx":
                        try:
                            import docx
                            doc = docx.Document(str(att_path))
                            paras = [p.text for p in doc.paragraphs if p.text.strip()]
                            for table in doc.tables:
                                for row in table.rows:
                                    r_text = " | ".join([c.text.strip() for c in row.cells if c.text.strip()])
                                    if r_text:
                                        paras.append(r_text)
                            doc_text = "\n\n".join(paras)
                            attachment_texts.append({
                                "filename": att_path.name,
                                "content": doc_text[:50000],
                            })
                            logger.info(f"Extracted DOCX text from {att_path.name} ({len(doc_text)} chars)")
                        except Exception as e:
                            logger.warning(f"Failed to read DOCX {att_path.name}: {e}")
                    elif ext == ".pdf":
                        try:
                            logger.info(f"Uploading PDF attachment to Gemini Files API: {att_path.name}")
                            pdf_file = self.client.files.upload(file=str(att_path))
                            uploaded_files.append(pdf_file)
                        except Exception as e:
                            logger.warning(f"Failed to upload PDF {att_path.name}: {e}")
                    else:
                        try:
                            content = att_path.read_text(encoding="utf-8", errors="replace")
                            attachment_texts.append({
                                "filename": att_path.name,
                                "content": content[:50000],
                            })
                            logger.info(f"Loaded text attachment {att_path.name} ({len(content)} chars)")
                        except Exception as e:
                            logger.warning(f"Failed to read attachment {att_path.name} as text: {e}")

            # 5. Construct user prompt with zero-hallucination guidance
            prompt_text = create_lecture_analysis_prompt(
                subject_hint=subject_hint,
                slides_info=slides_info,
                is_audio_silent=is_audio_silent,
                attachment_texts=attachment_texts,
                attachment_tables=extracted_tables,
                attachment_illustrations=attachment_illustrations,
            )

            # 5. Call Gemini Model with structured output schema
            contents = [*uploaded_files, prompt_text]

            # Candidate models: prioritize gemini-flash-latest, then specific flash versions
            standard_flash_models = [
                "gemini-flash-latest",
                "gemini-3.8-flash",
                "gemini-3.7-flash",
                "gemini-3.6-flash",
            ]
            if not self.model_name or self.model_name in standard_flash_models:
                candidate_models = list(standard_flash_models)
            else:
                candidate_models = [self.model_name] + [m for m in standard_flash_models if m != self.model_name]

            last_error = None
            response = None
            successful_model = None

            for model_cand in candidate_models:
                # 4 attempts per candidate model with backoff on 503 (patient retry on Google high demand spikes)
                max_attempts = 4
                for attempt in range(1, max_attempts + 1):
                    logger.info(f"Invoking Gemini model ({model_cand}) [attempt {attempt}/{max_attempts}]...")
                    try:
                        response = self.client.models.generate_content(
                            model=model_cand,
                            contents=contents,
                            config=types.GenerateContentConfig(
                                system_instruction=SYSTEM_INSTRUCTION,
                                response_mime_type="application/json",
                                response_schema=LectureAnalysisResult,
                                temperature=0.2,  # Low temperature for high factual accuracy
                                automatic_function_calling=types.AutomaticFunctionCallingConfig(disable=True),
                            ),
                        )
                        logger.info(f"Successfully generated analysis with model '{model_cand}'.")
                        successful_model = model_cand
                        break
                    except Exception as e:
                        last_error = e
                        err_str = str(e).lower()

                        # 1. Hard quota limit (e.g. limit: 0) -> skip immediately
                        if "limit: 0" in err_str or "limit:0" in err_str:
                            logger.warning(
                                f"Model '{model_cand}' has quota limit: 0. Skipping immediately to next candidate..."
                            )
                            break

                        # 2. Discontinued or missing model (404) -> skip immediately
                        if "404" in err_str or "not found" in err_str or "no longer available" in err_str:
                            logger.warning(f"Model '{model_cand}' is not available (404). Skipping immediately...")
                            break

                        # 3. Transient error (503 / 429 / high demand) -> retry up to 4 times
                        is_transient = any(
                            s in err_str
                            for s in [
                                "503", "429", "500", "unavailable",
                                "resource_exhausted", "high demand",
                                "spikes in demand", "try again later"
                            ]
                        )
                        if is_transient and attempt < max_attempts:
                            retry_delays = [20, 40, 60]
                            delay = retry_delays[attempt - 1] if (attempt - 1) < len(retry_delays) else 60
                            logger.warning(
                                f"Gemini API transient error on '{model_cand}' (attempt {attempt}/{max_attempts}): {e}. "
                                f"Retrying in {delay}s..."
                            )
                            import time
                            time.sleep(delay)
                        else:
                            logger.warning(
                                f"Attempt {attempt}/{max_attempts} failed on '{model_cand}': {e}. "
                                f"Failing over to next candidate model..."
                            )
                            break

                if response is not None:
                    break

            if response is None:
                logger.error(f"All Gemini models and retries failed. Last error: {last_error}")
                raise last_error

            # 5. Parse response and bind generating model name
            raw_text = response.text
            result = LectureAnalysisResult.model_validate_json(raw_text)
            actual_model = successful_model or self.model_name or "gemini-flash-latest"
            result.model_name = actual_model
            result.digest.model_name = actual_model
            result.longread.model_name = actual_model

            # 6. Sanitize slide references in sections to prevent broken image links
            self._sanitize_slide_references(result, effective_slide_paths)

            return result

        finally:
            # Clean up uploaded files in cloud to avoid storage accumulation
            for f in uploaded_files:
                try:
                    self.client.files.delete(name=f.name)
                except Exception as e:
                    logger.warning(f"Failed to delete cloud file {f.name}: {e}")

    def _sanitize_slide_references(
        self,
        result: LectureAnalysisResult,
        slide_paths: Optional[List[Union[str, Path]]],
    ):
        """
        Validates referenced_slide_filenames in each section against actual uploaded slide files.
        Removes hallucinated or non-existent slide references to prevent broken images in Markdown.
        """
        if not slide_paths or not result.longread or not result.longread.sections:
            return

        valid_names = {Path(p).name for p in slide_paths}
        # Also map clean stem prefixes (e.g. 'slide_001') to real filenames
        stem_map = {}
        for fn in valid_names:
            parts = fn.split("_")
            if len(parts) >= 2:
                prefix = f"{parts[0]}_{parts[1]}"  # e.g. slide_001
                stem_map[prefix] = fn
                stem_map[parts[0]] = fn

        for sec in result.longread.sections:
            cleaned_refs = []
            for ref in sec.referenced_slide_filenames:
                ref_clean = Path(ref).name
                if ref_clean in valid_names:
                    cleaned_refs.append(ref_clean)
                elif ref_clean in stem_map:
                    cleaned_refs.append(stem_map[ref_clean])
                else:
                    matched = next((fn for fn in valid_names if fn.startswith(Path(ref_clean).stem)), None)
                    if matched:
                        cleaned_refs.append(matched)
                    else:
                        logger.warning(
                            f"Filtered out invalid slide reference '{ref}' in section '{sec.heading}' (file not found in session)."
                        )
            sec.referenced_slide_filenames = cleaned_refs

    def _generate_simulated_result(
        self,
        subject_hint: str,
        slide_paths: Optional[List[Union[str, Path]]] = None,
        attachment_paths: Optional[List[Union[str, Path]]] = None,
    ) -> LectureAnalysisResult:
        """Generates realistic mock result for testing without external API."""
        subject = subject_hint or "Архитектура программных систем"
        slide_names = [Path(p).name for p in slide_paths] if slide_paths else ["slide_001_02m15s.png", "slide_002_18m40s.png"]
        att_names = [Path(p).name for p in attachment_paths] if attachment_paths else []

        digest = ActionItemsDigest(
            subject=subject,
            topic="Паттерны проектирования: Dependency Injection и Inversion of Control",
            date="Сегодня",
            executive_summary="Разобрали разницу между IoC и DI, типичные ошибки при внедрении зависимостей и архитектуру фабрик в современных фреймворках.",
            action_items=[
                ActionItem(
                    category="lab",
                    title="Лабораторная работа №3: Реализация DI-контейнера",
                    description="Написать свой легковесный DI-контейнер с поддержкой Singleton и Transient жизненных циклов. Сдавать строго в Git-репозитории группы.",
                    deadline="До следующей пятницы, 23:59",
                    importance="critical",
                ),
                ActionItem(
                    category="exam_hint",
                    title="Вопрос к зачету: Lifecycle зависимостей",
                    description="Преподаватель подчеркнул: обязательно спросит разницу между Scoped и Transient на реальном примере веб-сервера.",
                    importance="high",
                ),
            ],
            key_technologies_and_terms=[
                "Inversion of Control",
                "Dependency Injection",
                "Service Locator Anti-pattern",
                "Transient vs Singleton",
            ],
        )

        sections = [
            LectureSection(
                heading="Концепция Inversion of Control (IoC)",
                timecode="02m15s",
                content_markdown=(
                    "Традиционный подход к созданию зависимостей приводит к жесткой связанности (tight coupling):\n\n"
                    "```python\n"
                    "class OrderService:\n"
                    "    def __init__(self):\n"
                    "        # ПЛОХО: жесткая привязка к конкретной реализации\n"
                    "        self.notifier = EmailNotifier()\n"
                    "```\n\n"
                    "Принцип инверсии управления передает создание и связывание компонентов внешней сущности (контейнеру):\n\n"
                    "```python\n"
                    "class OrderService:\n"
                    "    def __init__(self, notifier: INotifier):\n"
                    "        # ХОРОШО: слабая связанность через интерфейс\n"
                    "        self.notifier = notifier\n"
                    "```"
                ),
                referenced_slide_filenames=[slide_names[0]] if slide_names else [],
                exam_cues=["Препод требует четко формулировать разницу между принципом (IoC) и конкретным паттерном (DI)."],
            ),
            LectureSection(
                heading="Жизненные циклы компонентов (Lifetimes)",
                timecode="18m40s",
                content_markdown=(
                    "Рассмотрены три ключевых жизненных цикла:\n"
                    "- **Transient**: новый экземпляр при каждом вызове.\n"
                    "- **Singleton**: один экземпляр на всё приложение.\n"
                    "- **Scoped**: один экземпляр на один HTTP-запрос / контекст сессии."
                ),
                referenced_slide_filenames=[slide_names[1]] if len(slide_names) > 1 else [],
                exam_cues=["Главная ловушка: внедрение Transient-сервиса в Singleton приводит к скрытым утечкам состояния."],
            ),
        ]

        longread = IllustratedLongread(
            title=f"{subject}: Dependency Injection",
            subject=subject,
            topic="Паттерны проектирования: Dependency Injection и Inversion of Control",
            overview="Конспект лекции, посвященной фундаментальным принципам слабой связанности в архитектуре программных систем.",
            sections=sections,
        )

        model_name = self.model_name or "gemini-flash-latest (simulation)"
        digest.model_name = model_name
        longread.model_name = model_name
        return LectureAnalysisResult(digest=digest, longread=longread, model_name=model_name)
