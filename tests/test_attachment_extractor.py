import io
import pytest
from pathlib import Path
from PIL import Image, ImageDraw
import imagehash

from ai.attachment_extractor import AttachmentExtractor, is_meaningful_image
from ai.prompts import create_lecture_analysis_prompt


def _create_test_image(width=400, height=300, color_variance=True) -> Image.Image:
    """Helper to create a test PIL image."""
    img = Image.new("RGB", (width, height), color=(255, 255, 255))
    if color_variance:
        draw = ImageDraw.Draw(img)
        # Draw some colored shapes for high variance
        draw.rectangle([10, 10, 150, 150], fill=(255, 0, 0))
        draw.rectangle([160, 10, 300, 150], fill=(0, 255, 0))
        draw.rectangle([10, 160, 300, 280], fill=(0, 0, 255))
        draw.text((20, 20), "Diagram Title", fill=(0, 0, 0))
    return img


def test_is_meaningful_image():
    # 1. Tiny icon
    tiny = Image.new("RGB", (32, 32), color=(100, 100, 100))
    assert not is_meaningful_image(tiny)

    # 2. Too small dimensions
    small = Image.new("RGB", (100, 150), color=(100, 100, 100))
    assert not is_meaningful_image(small)

    # 3. High aspect ratio (thin divider)
    divider = Image.new("RGB", (600, 30), color=(100, 100, 100))
    assert not is_meaningful_image(divider)

    # 4. Solid flat image (zero variance)
    flat = Image.new("RGB", (400, 300), color=(255, 255, 255))
    assert not is_meaningful_image(flat)

    # 5. Meaningful image with variance and proper size
    meaningful = _create_test_image(width=400, height=300, color_variance=True)
    assert is_meaningful_image(meaningful)


def test_attachment_extractor_standalone_image(tmp_path):
    dest_dir = tmp_path / "slides"
    dest_dir.mkdir()

    img_path = tmp_path / "custom_diagram.png"
    test_img = _create_test_image(500, 400, color_variance=True)
    test_img.save(img_path)

    extractor = AttachmentExtractor(output_dir=dest_dir)
    res = extractor.process_attachments([img_path], dest_slides_dir=dest_dir)

    assert len(res["extracted_images"]) == 1
    assert res["upgraded_slides_count"] == 0
    saved_file = res["extracted_images"][0]
    assert saved_file.exists()
    assert saved_file.name == "custom_diagram.png"
    assert len(res["illustrations_info"]) == 1
    assert res["illustrations_info"][0]["filename"] == "custom_diagram.png"


def test_attachment_extractor_phash_upgrade(tmp_path):
    slides_dir = tmp_path / "slides"
    slides_dir.mkdir()

    original_img = _create_test_image(800, 600, color_variance=True)

    # Simulate slightly blurry/downscaled stream slide
    blurry_slide = original_img.resize((320, 240)).resize((800, 600))
    slide_file = slides_dir / "slide_001_05m10s.png"
    blurry_slide.save(slide_file)

    # Original pristine slide attached by student
    pristine_file = tmp_path / "presentation_slide_1.png"
    original_img.save(pristine_file)

    extractor = AttachmentExtractor(output_dir=slides_dir)
    res = extractor.process_attachments(
        attachment_paths=[pristine_file],
        existing_slide_paths=[slide_file],
        dest_slides_dir=slides_dir,
    )

    # Should match via pHash and upgrade the blurry slide in place!
    assert res["upgraded_slides_count"] == 1
    assert len(res["extracted_images"]) == 0  # Not added as duplicate standalone

    # Check that slide_file was updated with pristine content
    with Image.open(slide_file) as updated:
        h1 = imagehash.phash(updated)
        h2 = imagehash.phash(original_img)
        assert h1 - h2 <= 2


def test_attachment_extractor_pptx_tables_and_images(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide_layout = prs.slide_layouts[5]  # Title only
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Архитектура Баз Данных"

    # Add a table
    rows, cols = 3, 2
    left, top, width, height = Inches(1), Inches(2), Inches(4), Inches(2)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.cell(0, 0).text = "СУБД"
    table.cell(0, 1).text = "Тип"
    table.cell(1, 0).text = "PostgreSQL"
    table.cell(1, 1).text = "Реляционная"
    table.cell(2, 0).text = "Redis"
    table.cell(2, 1).text = "Ключ-Значение"

    # Add an image to the slide
    img_data = _create_test_image(400, 300, color_variance=True)
    img_bytes = io.BytesIO()
    img_data.save(img_bytes, format="PNG")
    img_bytes.seek(0)
    slide.shapes.add_picture(img_bytes, Inches(1), Inches(4), Inches(3), Inches(2))

    pptx_path = tmp_path / "lecture.pptx"
    prs.save(str(pptx_path))

    slides_dir = tmp_path / "slides"
    extractor = AttachmentExtractor(output_dir=slides_dir)
    res = extractor.process_attachments([pptx_path], dest_slides_dir=slides_dir)

    # 1. Verify tables
    assert len(res["extracted_tables"]) == 1
    tbl = res["extracted_tables"][0]
    assert tbl["source_file"] == "lecture.pptx"
    assert "PostgreSQL" in tbl["markdown"]
    assert "Redis" in tbl["markdown"]
    assert "| :--- | :--- |" in tbl["markdown"]

    # 2. Verify slide text
    assert len(res["extracted_texts"]) == 1
    assert "Архитектура Баз Данных" in res["extracted_texts"][0]["content"]

    # 3. Verify extracted image
    assert len(res["extracted_images"]) >= 1
    assert any("pptx_lecture" in p.name for p in res["extracted_images"])


def test_attachment_extractor_pdf_extraction(tmp_path):
    import fitz

    doc = fitz.open()
    page = doc.new_page()

    # Add caption text (use ASCII so default Helvetica font renders without needing external TTF)
    page.insert_text((50, 50), "Fig. 1.1 - Client-Server Architecture Overview", fontsize=12)

    # Add an image
    test_img = _create_test_image(350, 250, color_variance=True)
    img_bytes = io.BytesIO()
    test_img.save(img_bytes, format="PNG")
    page.insert_image(fitz.Rect(50, 100, 400, 350), stream=img_bytes.getvalue())

    pdf_path = tmp_path / "manual.pdf"
    doc.save(str(pdf_path))
    doc.close()

    slides_dir = tmp_path / "slides"
    extractor = AttachmentExtractor(output_dir=slides_dir)
    res = extractor.process_attachments([pdf_path], dest_slides_dir=slides_dir)

    assert len(res["extracted_images"]) == 1
    ill_info = res["illustrations_info"][0]
    assert "manual.pdf" in ill_info["description"]
    assert "fig. 1.1" in ill_info["description"].lower()

    # Verify file descriptor is closed and file can be deleted immediately without Windows PermissionError
    pdf_path.unlink()
    assert not pdf_path.exists()


def test_attachment_extractor_pdf_handles_corrupt_file(tmp_path):
    bad_pdf = tmp_path / "corrupt.pdf"
    bad_pdf.write_bytes(b"not a valid pdf content")

    slides_dir = tmp_path / "slides"
    extractor = AttachmentExtractor(output_dir=slides_dir)
    res = extractor.process_attachments([bad_pdf], dest_slides_dir=slides_dir)
    assert res["extracted_images"] == []



def test_attachment_extractor_logo_buster(tmp_path):
    # Repeat the exact same image 3 times to simulate a template university logo
    slides_dir = tmp_path / "slides"
    logo_img = _create_test_image(300, 200, color_variance=True)

    files = []
    for i in range(3):
        p = tmp_path / f"logo_{i}.png"
        logo_img.save(p)
        files.append(p)

    extractor = AttachmentExtractor(output_dir=slides_dir)
    res = extractor.process_attachments(files, dest_slides_dir=slides_dir)

    # Because it appears >2 times, it should be recognized as a repeated template logo and ignored!
    assert len(res["extracted_images"]) == 0


def test_prompt_generation_with_tables_and_illustrations():
    tables = [
        {
            "source_file": "db.pptx",
            "title": "Сравнение СУБД",
            "markdown": "| СУБД | Тип |\n| :--- | :--- |\n| Postgres | SQL |",
        }
    ]
    illustrations = [
        {"filename": "att_arch.png", "description": "Диаграмма микросервисов"}
    ]

    prompt = create_lecture_analysis_prompt(
        subject_hint="Архитектура ПО",
        attachment_tables=tables,
        attachment_illustrations=illustrations,
    )

    assert "Сравнение СУБД" in prompt
    assert "| Postgres | SQL |" in prompt
    assert "att_arch.png" in prompt
    assert "Диаграмма микросервисов" in prompt


def test_illustrated_longread_format_markdown_attachments():
    from ai.schemas import IllustratedLongread, LectureSection

    sec1 = LectureSection(
        heading="Раздел 1",
        timecode="00:00",
        content_markdown="Описание темы",
        referenced_slide_filenames=["slide_001_00m00s.png", "IOS_contexts.png"],
    )
    longread = IllustratedLongread(
        title="Тест",
        subject="Сети",
        topic="Маршрутизация",
        overview="Обзор",
        sections=[sec1],
    )

    md = longread.format_markdown(assets_relative_path="slides")
    assert "![slide_001_00m00s.png](slides/slide_001_00m00s.png)" in md
    assert "![IOS_contexts.png](attachments/IOS_contexts.png)" in md
